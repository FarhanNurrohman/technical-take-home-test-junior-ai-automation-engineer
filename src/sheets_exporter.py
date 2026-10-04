"""Build and export settlement reports to Google Sheets."""

from __future__ import annotations

import json
import logging
import math
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd

from src import config

logger = logging.getLogger(__name__)
_RESULT_COLUMNS = 12
_UNMATCHED_COLUMNS = 7
_DASHBOARD_COLUMNS = 8
_EXPECTED_SHEETS = ("Working_Paper_Result", "Dashboard", "Unmatched_GL")
_STATUS_COLORS = {
	"PARTIAL": "#FFF2CC",
	"UNSETTLED": "#FFF2CC",
	"OVER_SETTLED": "#F4CCCC",
	"SETTLED": "#D9EAD3",
}


def _a1_range_bounds(cell_range: str) -> tuple[int, int, int, int]:
	"""Convert a simple A1 range into zero-based half-open row/column bounds."""
	match = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", cell_range.upper())
	if not match:
		raise ValueError(f"Format merge tidak valid: {cell_range}")
	start_column = _column_index(match.group(1))
	end_column = _column_index(match.group(3)) + 1
	start_row = int(match.group(2)) - 1
	end_row = int(match.group(4))
	if start_row >= end_row or start_column >= end_column:
		raise ValueError(f"Range merge kosong atau terbalik: {cell_range}")
	return start_row, end_row, start_column, end_column


def validate_payload(payload: dict[str, Any]) -> None:
	"""Validate that frozen boundaries do not split any merged cell."""
	freeze = payload.get("freeze", {})
	frozen_rows = int(freeze.get("rows", 0))
	frozen_columns = int(freeze.get("columns", 0))
	if frozen_rows < 0 or frozen_columns < 0:
		raise ValueError("Batas freeze tidak boleh negatif.")
	for merge in payload.get("merges", []):
		cell_range = str(merge.get("range", ""))
		start_row, end_row, start_column, end_column = _a1_range_bounds(cell_range)
		if start_row < frozen_rows < end_row or start_column < frozen_columns < end_column:
			raise ValueError(
				f"Merge {cell_range} berpotongan dengan batas freeze "
				f"(rows={frozen_rows}, columns={frozen_columns})."
			)


def _missing(value: Any) -> bool:
	if value is None:
		return True
	try:
		return bool(pd.isna(value))
	except (TypeError, ValueError):
		return False


def _number(value: Any) -> int | float:
	if _missing(value):
		return 0
	try:
		number = float(value)
	except (TypeError, ValueError):
		return 0
	if not math.isfinite(number):
		return 0
	return int(number) if number.is_integer() else number


def _text(value: Any) -> str:
	if _missing(value):
		return ""
	return str(value)


def _safe_text(value: Any) -> str:
	text = _text(value)
	if text.startswith(("=", "+", "-", "@")):
		return "'" + text
	return text


def _date_value(value: Any) -> str:
	if _missing(value):
		return ""
	try:
		converted = pd.to_datetime(value, errors="coerce")
		if pd.isna(converted):
			return _safe_text(value)
		return converted.date().isoformat()
	except (TypeError, ValueError, OverflowError):
		return _safe_text(value)


def _records(frame: pd.DataFrame | None) -> list[dict[str, Any]]:
	if frame is None or frame.empty:
		return []
	return frame.to_dict(orient="records")


def _value(row: dict[str, Any], *keys: str, default: Any = None) -> Any:
	for key in keys:
		if key in row and not _missing(row[key]):
			return row[key]
	return default


def _blank_row(width: int) -> list[Any]:
	return [""] * width


def _add_format(
	payload: dict[str, Any],
	cell_range: str,
	*,
	background: str | None = None,
	bold: bool = False,
	wrap: bool = False,
	align: str | None = None,
	number_format: str | None = None,
	font_color: str | None = None,
) -> None:
	format_spec: dict[str, Any] = {"range": cell_range}
	if background:
		format_spec["background"] = background
	if bold:
		format_spec["bold"] = True
	if wrap:
		format_spec["wrap"] = True
	if align:
		format_spec["align"] = align
	if number_format:
		format_spec["number_format"] = number_format
	if font_color:
		format_spec["font_color"] = font_color
	payload["formats"].append(format_spec)


def build_result_payload(df_result: pd.DataFrame) -> dict[str, Any]:
	"""Build Working_Paper_Result values and formatting without changing input data."""
	rows = _records(df_result)
	wp_rows = sorted((row for row in rows if _text(row.get("section", "WP")) == "WP"), key=lambda row: _number(row.get("wp_row")))
	new_rows = sorted((row for row in rows if _text(row.get("section")) == "NEW_ADVANCE"), key=lambda row: _date_value(row.get("date")))
	values: list[list[Any]] = [
		["PT. SETIAWAN DWI TUNGGAL"] + [""] * (_RESULT_COLUMNS - 1),
		["Uang Muka Pihak Ketiga Lainnya"] + [""] * (_RESULT_COLUMNS - 1),
		["Periode: April 2026"] + [""] * (_RESULT_COLUMNS - 1),
		["110.040.040.000 Advances - Other"] + [""] * (_RESULT_COLUMNS - 1),
		_blank_row(_RESULT_COLUMNS),
		["Date", "Voucher No", "Description", "Amount", "Realization", "", "", "Saldo", "Description (catatan)", "Status", "Flag", "Target ID"],
		["", "", "", "", "Date", "No. Voucher", "Amount", "", "", "", "", ""],
	]
	meta: dict[str, Any] = {"section_a_rows": [], "section_b_rows": []}
	formats: list[dict[str, Any]] = []
	merges = [{"range": "E6:G6"}, {"range": "A1:L1"}, {"range": "A2:L2"}, {"range": "A3:L3"}, {"range": "A4:L4"}]
	values.append(["A. ADVANCE AWAL (Working Paper)"] + [""] * (_RESULT_COLUMNS - 1))
	section_a_label = len(values)
	_add_format({"formats": formats}, f"A{section_a_label}:L{section_a_label}", background="#D9E1F2", bold=True)

	def append_data(row: dict[str, Any], section: str) -> None:
		row_number = len(values) + 1
		realization = _number(_value(row, "realization_amount", "realized", default=0))
		amount = _number(row.get("amount"))
		voucher = _value(row, "realization_vouchers", "realization_voucher_no", "realization_voucher", default="")
		flag_parts = []
		if bool(_value(row, "need_settlement_evidence", default=False)):
			flag_parts.append("BUTUH BUKTI REALISASI")
		if _text(row.get("amount_check")).upper() == "DIFF":
			flag_parts.append("AMOUNT DIFF")
		warning = _text(row.get("warning"))
		if warning:
			flag_parts.append(warning)
		if _text(row.get("status")).upper() == "OVER_SETTLED":
			flag_parts.append("OVER SETTLED")
		values.append(
			[
				_date_value(row.get("date")),
				_safe_text(row.get("voucher_no")),
				_safe_text(row.get("description")),
				amount,
				_date_value(_value(row, "realization_date", default="")),
				_safe_text(voucher),
				realization,
				f"=D{row_number}-G{row_number}",
				_safe_text(row.get("note")),
				_safe_text(row.get("status")),
				_safe_text("; ".join(flag_parts)),
				_safe_text(row.get("target_id")),
			]
		)
		meta[f"section_{'a' if section == 'WP' else 'b'}_rows"].append(row_number)
		color = _STATUS_COLORS.get(_text(row.get("status")).upper())
		if color:
			_add_format({"formats": formats}, f"A{row_number}:L{row_number}", background=color, wrap=True)
		_add_format({"formats": formats}, f"A{row_number}:A{row_number}", number_format="dd-mmm-yyyy")
		_add_format({"formats": formats}, f"E{row_number}:E{row_number}", number_format="dd-mmm-yyyy")
		_add_format({"formats": formats}, f"D{row_number}:D{row_number}", number_format="#,##0")
		_add_format({"formats": formats}, f"G{row_number}:H{row_number}", number_format="#,##0")
		if flag_parts:
			_add_format({"formats": formats}, f"K{row_number}:K{row_number}", bold=True, wrap=True)

	for row in wp_rows:
		append_data(row, "WP")
	meta["section_a_start"] = meta["section_a_rows"][0] if meta["section_a_rows"] else None
	meta["section_a_end"] = meta["section_a_rows"][-1] if meta["section_a_rows"] else None
	meta["subtotal_a_row"] = len(values) + 1
	values.append(["Subtotal A"] + [""] * 2 + [_sum_formula("D", meta["section_a_rows"])] + [""] * 2 + [_sum_formula("G", meta["section_a_rows"]), _sum_formula("H", meta["section_a_rows"])] + [""] * 4)
	_add_format({"formats": formats}, f"A{meta['subtotal_a_row']}:L{meta['subtotal_a_row']}", background="#D9D9D9", bold=True)

	values.append(_blank_row(_RESULT_COLUMNS))
	values.append(["B. ADVANCE BARU APRIL 2026 (dari GL DEBET)"] + [""] * (_RESULT_COLUMNS - 1))
	section_b_label = len(values)
	_add_format({"formats": formats}, f"A{section_b_label}:L{section_b_label}", background="#D9E1F2", bold=True)
	if new_rows:
		for row in new_rows:
			append_data(row, "NEW_ADVANCE")
	else:
		values.append(["Tidak ada advance baru"] + [""] * (_RESULT_COLUMNS - 1))
	meta["section_b_start"] = meta["section_b_rows"][0] if meta["section_b_rows"] else None
	meta["section_b_end"] = meta["section_b_rows"][-1] if meta["section_b_rows"] else None
	meta["subtotal_b_row"] = len(values) + 1
	if new_rows:
		b_amount = _sum_formula("D", meta["section_b_rows"])
		b_realization = _sum_formula("G", meta["section_b_rows"])
		b_balance = _sum_formula("H", meta["section_b_rows"])
	else:
		b_amount = b_realization = b_balance = 0
	values.append(["Subtotal B"] + [""] * 2 + [b_amount] + [""] * 2 + [b_realization, b_balance] + [""] * 4)
	_add_format({"formats": formats}, f"A{meta['subtotal_b_row']}:L{meta['subtotal_b_row']}", background="#D9D9D9", bold=True)
	meta["grand_total_row"] = len(values) + 1
	values.append(["GRAND TOTAL (A + B; Total Advance Awal = Subtotal A)", "", "", f"=D{meta['subtotal_a_row']}+D{meta['subtotal_b_row']}", "", "", f"=G{meta['subtotal_a_row']}+G{meta['subtotal_b_row']}", f"=H{meta['subtotal_a_row']}+H{meta['subtotal_b_row']}", "", "", "", ""])
	_add_format({"formats": formats}, f"A{meta['grand_total_row']}:L{meta['grand_total_row']}", background="#B7C9E2", bold=True)
	meta["values_rows"] = len(values)
	_add_format({"formats": formats}, f"A8:L{meta['grand_total_row']}", wrap=True)
	_add_format({"formats": formats}, "A6:L7", background="#1F4E78", bold=True, wrap=True, align="CENTER")
	_add_format({"formats": formats}, "A6:L7", font_color="#FFFFFF")
	_add_format({"formats": formats}, "A1:L1", bold=True)
	return {
		"values": values,
		"formats": formats,
		"merges": merges,
		"column_widths": [{"column": "A", "width": 110}, {"column": "B", "width": 190}, {"column": "C", "width": 360}, {"column": "D", "width": 125}, {"column": "E", "width": 125}, {"column": "F", "width": 300}, {"column": "G", "width": 130}, {"column": "H", "width": 130}, {"column": "I", "width": 300}, {"column": "J", "width": 135}, {"column": "K", "width": 220}, {"column": "L", "width": 200}],
		"freeze": {"rows": 7, "columns": 0},
		"meta": meta,
	}


def _sum_formula(column: str, rows: list[int]) -> str:
	if not rows:
		return "=0"
	return f"=SUM({column}{rows[0]}:{column}{rows[-1]})"


def build_unmatched_payload(
	df_unmatched: pd.DataFrame | None,
	df_debit_exceptions: pd.DataFrame | None = None,
	df_suggestions: pd.DataFrame | None = None,
	summarize_sections: Any = None,
) -> dict[str, Any]:
	"""Build the GL audit sheet, including exceptions and manual suggestions."""
	values: list[list[Any]] = [["KREDIT/ADJUSTMENT TIDAK TER-MATCH"] + [""] * (_UNMATCHED_COLUMNS - 1)]
	formats: list[dict[str, Any]] = []
	merges: list[dict[str, Any]] = []
	_add_format({"formats": formats}, "A1:G1", background="#D9D9D9", bold=True)
	unmatched = _records(df_unmatched)
	credit_rows = [row for row in unmatched if _text(row.get("txn_type")).upper() in {"", "SETTLEMENT", "REFUND", "ADJUSTMENT"}]
	values.append(["GL row", "Tanggal", "No Jurnal", "Txn Type", "Nominal", "Alasan", "Deskripsi"])
	_add_format({"formats": formats}, "A2:G2", background="#1F4E78", bold=True, wrap=True, font_color="#FFFFFF")
	for row in credit_rows:
		amount_keys = ("debit", "amount", "credit") if _text(row.get("txn_type")).upper() == "ADJUSTMENT" else ("credit", "amount", "debit")
		amount = _number(_value(row, *amount_keys, default=0))
		values.append([
			_number(row.get("gl_row")),
			_date_value(_value(row, "gl_date", "date")),
			_safe_text(_value(row, "voucher_no", "journal_no", default="")),
			_safe_text(row.get("txn_type")),
			amount,
			_safe_text(row.get("reason")),
			_safe_text(row.get("description")),
		])
		row_number = len(values)
		_add_format({"formats": formats}, f"B{row_number}:B{row_number}", number_format="dd-mmm-yyyy")
		_add_format({"formats": formats}, f"E{row_number}:E{row_number}", number_format="#,##0")
	if not credit_rows:
		values.append(["Semua transaksi KREDIT ter-match"] + [""] * (_UNMATCHED_COLUMNS - 1))
	credit_total_row = len(values) + 1
	values.append(["Total nominal KREDIT/ADJUSTMENT unmatched", "", "", "", f"=SUM(E3:E{credit_total_row - 1})" if credit_rows else 0, "", ""])
	_add_format({"formats": formats}, f"A{credit_total_row}:G{credit_total_row}", bold=True, background="#EDEDED")

	values.append(_blank_row(_UNMATCHED_COLUMNS))
	debit_title_row = len(values) + 1
	values.append(["DEBET PERLU DICEK"] + [""] * (_UNMATCHED_COLUMNS - 1))
	_add_format({"formats": formats}, f"A{debit_title_row}:G{debit_title_row}", background="#D9D9D9", bold=True)
	values.append(["GL row", "Tanggal", "No Jurnal", "Txn Type", "Nominal", "Alasan", "Deskripsi"])
	debit_header_row = len(values)
	_add_format({"formats": formats}, f"A{debit_header_row}:G{debit_header_row}", background="#1F4E78", bold=True, wrap=True, font_color="#FFFFFF")
	for row in _records(df_debit_exceptions):
		values.append([_number(row.get("gl_row")), _date_value(_value(row, "gl_date", "date")), _safe_text(_value(row, "voucher_no", "journal_no", default="")), _safe_text(_value(row, "txn_type", default="UNCLASSIFIED_DEBIT")), _number(_value(row, "amount", "debit", default=0)), _safe_text(row.get("reason")), _safe_text(row.get("description"))])
		row_number = len(values)
		_add_format({"formats": formats}, f"B{row_number}:B{row_number}", number_format="dd-mmm-yyyy")
		_add_format({"formats": formats}, f"E{row_number}:E{row_number}", number_format="#,##0")

	values.append(_blank_row(_UNMATCHED_COLUMNS))
	suggestion_title_row = len(values) + 1
	values.append(["SARAN SETTLEMENT (bukan match otomatis)"] + [""] * (_UNMATCHED_COLUMNS - 1))
	_add_format({"formats": formats}, f"A{suggestion_title_row}:G{suggestion_title_row}", background="#D9D9D9", bold=True)
	values.append(["Target", "Baris GL", "Voucher", "Nominal", "Alasan", "", ""])
	suggestion_header_row = len(values)
	_add_format({"formats": formats}, f"A{suggestion_header_row}:G{suggestion_header_row}", background="#1F4E78", bold=True, wrap=True, font_color="#FFFFFF")
	for row in _records(df_suggestions):
		values.append([_safe_text(_value(row, "target", "target_id", default="")), _number(_value(row, "gl_row", "row", default=0)), _safe_text(_value(row, "voucher_no", "journal_no", default="")), _number(_value(row, "amount", "credit", default=0)), _safe_text(row.get("reason")), "", ""])
	if not _records(df_suggestions):
		values.append(["SARAN perlu diverifikasi manual; bukan match otomatis"] + [""] * (_UNMATCHED_COLUMNS - 1))
	else:
		values.append(["SARAN perlu diverifikasi manual; bukan match otomatis"] + [""] * (_UNMATCHED_COLUMNS - 1))
	_add_format({"formats": formats}, f"A{len(values)}:G{len(values)}", bold=True, wrap=True)
	return {
		"values": values,
		"formats": formats,
		"merges": merges,
		"column_widths": [{"column": "A", "width": 190}, {"column": "B", "width": 105}, {"column": "C", "width": 200}, {"column": "D", "width": 150}, {"column": "E", "width": 130}, {"column": "F", "width": 180}, {"column": "G", "width": 420}],
		"freeze": {"rows": 2, "columns": 0},
		"meta": {"credit_total_row": credit_total_row, "unmatched_count": len(credit_rows), "debit_title_row": debit_title_row, "suggestion_title_row": suggestion_title_row},
	}


def _plain_summary(summary_text: str) -> list[str]:
	lines: list[str] = []
	for original in str(summary_text or "").splitlines():
		line = original.strip()
		if not line:
			continue
		line = re.sub(r"^\s*#{1,6}\s*", "", line)
		line = line.replace("**", "").replace("__", "").replace("`", "")
		line = re.sub(r"^[-*+]\s+", "• ", line)
		lines.append(_safe_text(line))
	return lines or [""]


def build_dashboard_payload(
	metrics: dict[str, Any],
	summary_text: str,
	df_result: pd.DataFrame,
	result_meta: dict[str, Any],
) -> dict[str, Any]:
	"""Build KPI, quality, summary, and open-item sections for Dashboard."""
	wp = metrics.get("wp", {}) or {}
	values: list[list[Any]] = [
		["Dashboard Advance Settlement April 2026"] + [""] * (_DASHBOARD_COLUMNS - 1),
		_blank_row(_DASHBOARD_COLUMNS),
		["KPI", "Advance Awal (WP)", "Advance Baru April", "Total", "", "", "", ""],
		["Total Advance", f"='Working_Paper_Result'!D{result_meta['subtotal_a_row']}", f"='Working_Paper_Result'!D{result_meta['subtotal_b_row']}", "=B4+C4", "", "", "", ""],
		["Total Realisasi (G)", f"='Working_Paper_Result'!G{result_meta['subtotal_a_row']}", f"='Working_Paper_Result'!G{result_meta['subtotal_b_row']}", "=B5+C5", "", "", "", ""],
		["Total Sisa Saldo (H)", f"='Working_Paper_Result'!H{result_meta['subtotal_a_row']}", f"='Working_Paper_Result'!H{result_meta['subtotal_b_row']}", "=B6+C6", "", "", "", ""],
		["Jumlah Item Unsettled/Partial", _countif_formula("A", result_meta), _countif_formula("B", result_meta), "=B7+C7", "", "", "", ""],
		["Total Advance Awal = hanya baris Working Paper asli"] + [""] * (_DASHBOARD_COLUMNS - 1),
		_blank_row(_DASHBOARD_COLUMNS),
		["Kualitas Data", "Jumlah / Nilai"] + [""] * (_DASHBOARD_COLUMNS - 2),
		["Kredit GL unmatched", _number(metrics.get("unmatched_count"))] + [""] * (_DASHBOARD_COLUMNS - 2),
		["Total kredit GL unmatched", _number(metrics.get("unmatched_total"))] + [""] * (_DASHBOARD_COLUMNS - 2),
		["Target butuh bukti realisasi", _number(metrics.get("need_settlement_evidence_count"))] + [""] * (_DASHBOARD_COLUMNS - 2),
		["Target OVER_SETTLED", _number(metrics.get("over_settled_count"))] + [""] * (_DASHBOARD_COLUMNS - 2),
		["Rekonsiliasi", "Selisih = 0" if bool(metrics.get("reconcile_ok")) else f"Tidak seimbang ({_number(metrics.get('reconcile_diff'))})"] + [""] * (_DASHBOARD_COLUMNS - 2),
		_blank_row(_DASHBOARD_COLUMNS),
	]
	formats: list[dict[str, Any]] = []
	merges = [{"range": "A1:H1"}]
	row_heights: list[dict[str, int]] = []
	_add_format({"formats": formats}, "A3:D3", background="#1F4E78", bold=True, wrap=True)
	_add_format({"formats": formats}, "A1:H1", background="#1F4E78", bold=True, font_color="#FFFFFF")
	_add_format({"formats": formats}, "A3:D3", font_color="#FFFFFF")
	_add_format({"formats": formats}, "A10:B10", background="#D9D9D9", bold=True)
	summary_title_row = len(values) + 1
	values.append(["Executive Summary"] + [""] * (_DASHBOARD_COLUMNS - 1))
	row_heights.append({"row": summary_title_row, "height": 24})
	merges.append({"range": f"A{summary_title_row}:F{summary_title_row}", "purpose": "summary"})
	_add_format({"formats": formats}, f"A{summary_title_row}:F{summary_title_row}", bold=True, background="#D9E1F2")
	for row_text in _plain_summary(summary_text):
		row_number = len(values) + 1
		values.append([row_text] + [""] * (_DASHBOARD_COLUMNS - 1))
		row_heights.append({"row": row_number, "height": max(21, 18 * math.ceil(len(row_text) / 140))})
		merges.append({"range": f"A{row_number}:F{row_number}", "purpose": "summary"})
		_add_format({"formats": formats}, f"A{row_number}:F{row_number}", wrap=True)
	values.append(_blank_row(_DASHBOARD_COLUMNS))
	detail_title_row = len(values) + 1
	values.append(["Rincian Item Unsettled / Partial"] + [""] * (_DASHBOARD_COLUMNS - 1))
	_add_format({"formats": formats}, f"A{detail_title_row}:H{detail_title_row}", background="#D9E1F2", bold=True)
	values.append(["Bagian", "Deskripsi", "Amount", "Realisasi", "Saldo", "Status", "Flag", "Saran"])
	_add_format({"formats": formats}, f"A{detail_title_row + 1}:H{detail_title_row + 1}", background="#1F4E78", bold=True, wrap=True)
	_add_format({"formats": formats}, f"A{detail_title_row + 1}:H{detail_title_row + 1}", font_color="#FFFFFF")
	items = metrics.get("unsettled_items", []) or []
	items = sorted(items, key=lambda item: _number(item.get("balance")), reverse=True)
	for item in items:
		values.append([
			_safe_text(item.get("section")),
			_safe_text(item.get("description")),
			_number(item.get("amount")),
			_number(_value(item, "realized", "realization_amount", default=0)),
			_number(_value(item, "balance", "saldo", default=0)),
			_safe_text(item.get("status")),
			_safe_text(", ".join(item.get("flags", [])) if isinstance(item.get("flags"), list) else item.get("flags")),
			_safe_text(_value(item, "suggested_action", "suggestion", default="")),
		])
		_add_format({"formats": formats}, f"A{len(values)}:H{len(values)}", wrap=True)
	values.append([f"Diperbarui: {datetime.now().astimezone().isoformat(timespec='seconds')}"] + [""] * (_DASHBOARD_COLUMNS - 1))
	return {
		"values": values,
		"formats": formats,
		"merges": merges,
		"column_widths": [{"column": "A", "width": 190}, {"column": "B", "width": 330}, {"column": "C", "width": 150}, {"column": "D", "width": 150}, {"column": "E", "width": 150}, {"column": "F", "width": 150}, {"column": "G", "width": 220}, {"column": "H", "width": 430}],
		"freeze": {"rows": 3, "columns": 0},
		"row_heights": row_heights,
		"meta": {"summary_title_row": summary_title_row, "detail_title_row": detail_title_row, "result_meta": result_meta},
	}


def _countif_formula(section: str, meta: dict[str, Any]) -> str:
	start = meta.get(f"section_{section.lower()}_start")
	end = meta.get(f"section_{section.lower()}_end")
	if start is None or end is None:
		return "=0"
	return f"=COUNTIF('Working_Paper_Result'!H{start}:H{end},\">0.01\")"


def _column_index(column: str) -> int:
	index = 0
	for char in column.upper():
		index = index * 26 + ord(char) - ord("A") + 1
	return index - 1


def _cell_index(cell: str) -> tuple[int, int]:
	match = re.fullmatch(r"([A-Z]+)(\d+)", cell.upper())
	if not match:
		raise ValueError(f"Rentang Google Sheets tidak valid: {cell}")
	return int(match.group(2)) - 1, _column_index(match.group(1))


def _grid_range(cell_range: str, sheet_id: int) -> dict[str, int]:
	start, end = cell_range.split(":", maxsplit=1)
	start_row, start_col = _cell_index(start)
	end_row, end_col = _cell_index(end)
	return {"sheetId": sheet_id, "startRowIndex": start_row, "endRowIndex": end_row + 1, "startColumnIndex": start_col, "endColumnIndex": end_col + 1}


def _color(hex_color: str) -> dict[str, float]:
	clean = hex_color.lstrip("#")
	return {"red": int(clean[0:2], 16) / 255, "green": int(clean[2:4], 16) / 255, "blue": int(clean[4:6], 16) / 255}


def _format_request(spec: dict[str, Any], sheet_id: int) -> dict[str, Any]:
	format_body: dict[str, Any] = {}
	fields: list[str] = []
	if spec.get("background"):
		format_body["backgroundColor"] = _color(spec["background"])
		fields.append("userEnteredFormat.backgroundColor")
	if spec.get("bold"):
		format_body.setdefault("textFormat", {})["bold"] = True
		fields.append("userEnteredFormat.textFormat.bold")
	if spec.get("font_color"):
		format_body.setdefault("textFormat", {})["foregroundColor"] = _color(spec["font_color"])
		fields.append("userEnteredFormat.textFormat.foregroundColor")
	if spec.get("wrap"):
		format_body["wrapStrategy"] = "WRAP"
		fields.append("userEnteredFormat.wrapStrategy")
	if spec.get("align"):
		format_body["horizontalAlignment"] = spec["align"]
		fields.append("userEnteredFormat.horizontalAlignment")
	if spec.get("number_format"):
		pattern = spec["number_format"]
		format_body["numberFormat"] = {"type": "DATE", "pattern": pattern} if "mmm" in pattern else {"type": "NUMBER", "pattern": pattern}
		fields.append("userEnteredFormat.numberFormat")
	return {"repeatCell": {"range": _grid_range(spec["range"], sheet_id), "cell": {"userEnteredFormat": format_body}, "fields": ",".join(fields)}}


def _payload_requests(payload: dict[str, Any], sheet_id: int, rows: int, columns: int) -> list[dict[str, Any]]:
	requests: list[dict[str, Any]] = [
		{"unmergeCells": {"range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": rows, "startColumnIndex": 0, "endColumnIndex": columns}}},
		{"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": rows, "startColumnIndex": 0, "endColumnIndex": columns}, "cell": {"userEnteredFormat": {}}, "fields": "userEnteredFormat"}},
	]
	for spec in payload["formats"]:
		requests.append(_format_request(spec, sheet_id))
	for merge in payload["merges"]:
		requests.append({"mergeCells": {"range": _grid_range(merge["range"], sheet_id), "mergeType": "MERGE_ALL"}})
	for width in payload["column_widths"]:
		column_index = _column_index(width["column"])
		requests.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS", "startIndex": column_index, "endIndex": column_index + 1}, "properties": {"pixelSize": width["width"]}, "fields": "pixelSize"}})
	for height in payload.get("row_heights", []):
		row_index = height["row"] - 1
		requests.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "ROWS", "startIndex": row_index, "endIndex": row_index + 1}, "properties": {"pixelSize": height["height"]}, "fields": "pixelSize"}})
	requests.append({"updateSheetProperties": {"properties": {"sheetId": sheet_id, "gridProperties": {"frozenRowCount": payload["freeze"]["rows"], "frozenColumnCount": payload["freeze"]["columns"]}}, "fields": "gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}})
	return requests


def _is_quota_error(error: Exception) -> bool:
	status = getattr(getattr(error, "response", None), "status_code", None)
	status = status or getattr(error, "code", None)
	return status == 429 or "429" in str(error) or "quota" in str(error).lower()


def _run_with_retry(operation: Any) -> Any:
	for attempt in range(3):
		try:
			return operation()
		except Exception as error:
			if not _is_quota_error(error) or attempt == 2:
				raise
			logger.warning("Google Sheets quota response; retry %s/2", attempt + 1)
			time.sleep(2**attempt)
	raise RuntimeError("Google Sheets operation ended unexpectedly")


def _create_client(*, read_only: bool = False) -> tuple[Any, str | None]:
	credential_path = config.GOOGLE_SHEETS_CRED
	if not credential_path:
		raise ValueError("GOOGLE_SHEETS_CRED belum dikonfigurasi di environment.")
	path = Path(credential_path).expanduser()
	if not path.is_file():
		raise FileNotFoundError(f"File kredensial Google Sheets tidak ditemukan: {path}")
	try:
		with path.open("r", encoding="utf-8") as credential_file:
			credential_data = json.load(credential_file)
	except (OSError, json.JSONDecodeError) as error:
		raise ValueError(f"File kredensial Google Sheets tidak dapat dibaca sebagai JSON valid: {path}") from error
	service_email = credential_data.get("client_email") if isinstance(credential_data, dict) else None
	try:
		import gspread
		from google.oauth2.service_account import Credentials

		scopes = (
			["https://www.googleapis.com/auth/spreadsheets.readonly"]
			if read_only
			else ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
		)
		credentials = Credentials.from_service_account_file(
			str(path), scopes=scopes
		)
		return gspread.authorize(credentials), service_email
	except (ImportError, ValueError, OSError) as error:
		logger.error("Google Sheets client initialization failed; verify credentials and dependencies")
		raise RuntimeError("Gagal menginisialisasi autentikasi Google Sheets; periksa file kredensial dan dependensi.") from error


def create_sheets_client(*, read_only: bool = False) -> Any:
	"""Create an authenticated gspread client for report access."""
	client, _ = _create_client(read_only=read_only)
	return client


def _ensure_worksheets(spreadsheet: Any) -> dict[str, Any]:
	worksheets = {worksheet.title: worksheet for worksheet in spreadsheet.worksheets()}
	for title in _EXPECTED_SHEETS:
		if title not in worksheets:
			worksheets[title] = spreadsheet.add_worksheet(title=title, rows=100, cols=12)
	if "Sheet1" in worksheets and len(worksheets) > len(_EXPECTED_SHEETS):
		spreadsheet.del_worksheet(worksheets["Sheet1"])
		worksheets.pop("Sheet1", None)
	ordered = [worksheets[title] for title in _EXPECTED_SHEETS]
	spreadsheet.reorder_worksheets(ordered)
	return worksheets


def _sheet_id(worksheet: Any, index: int) -> int:
	return int(getattr(worksheet, "id", index))


def _write_payload(spreadsheet: Any, worksheet: Any, payload: dict[str, Any], index: int) -> None:
	validate_payload(payload)
	rows = max(1, len(payload["values"]))
	columns = max((len(row) for row in payload["values"]), default=1)
	grid_rows = max(rows, int(getattr(worksheet, "row_count", rows)))
	grid_columns = max(columns, int(getattr(worksheet, "col_count", columns)))
	_run_with_retry(worksheet.clear)
	_run_with_retry(lambda: worksheet.resize(rows=grid_rows, cols=grid_columns))
	_run_with_retry(lambda: worksheet.update("A1", payload["values"], value_input_option="USER_ENTERED"))
	sheet_id = _sheet_id(worksheet, index)
	requests = _payload_requests(payload, sheet_id, grid_rows, grid_columns)
	if hasattr(spreadsheet, "fetch_sheet_metadata"):
		metadata = spreadsheet.fetch_sheet_metadata()
		for sheet in metadata.get("sheets", []):
			if sheet.get("properties", {}).get("sheetId") == sheet_id:
				for rule_index in reversed(range(len(sheet.get("conditionalFormats", [])))):
					requests.append({"deleteConditionalFormatRule": {"sheetId": sheet_id, "index": rule_index}})
				break
	try:
		_run_with_retry(lambda: spreadsheet.batch_update({"requests": requests}))
	except Exception as error:
		request_match = re.search(r"requests\[(\d+)\]", str(error))
		request_index = request_match.group(1) if request_match else "unknown"
		raise RuntimeError(
			f"Google Sheets batch_update gagal untuk worksheet '{worksheet.title}', "
			f"request [{request_index}]: {error}"
		) from error
	if grid_rows != rows or grid_columns != columns:
		_run_with_retry(lambda: worksheet.resize(rows=rows, cols=columns))


def _raise_google_error(error: Exception, service_email: str | None, client: Any) -> None:
	message = str(error).lower()
	if "google sheets batch_update gagal untuk worksheet" in message:
		raise RuntimeError(str(error)) from error
	if "403" in message or "permission" in message or "not found" in message:
		email = service_email or getattr(getattr(client, "auth", None), "service_account_email", None) or "email service account"
		raise PermissionError(f"Spreadsheet tidak dapat diakses. Bagikan spreadsheet ke {email} dengan akses Editor.") from error
	if "api" in message and ("disabled" in message or "not been used" in message):
		raise RuntimeError("Google Sheets API atau Google Drive API belum aktif di Google Cloud project.") from error
	if _is_quota_error(error):
		raise RuntimeError("Kuota Google Sheets habis setelah 2 kali retry.") from error
	raise RuntimeError("Ekspor gagal; periksa izin spreadsheet, status Sheets/Drive API, dan kuota Google.") from error


def export_to_sheets(
	df_result: pd.DataFrame,
	summary_text: str,
	metrics: dict[str, Any],
	df_unmatched: pd.DataFrame,
	df_debit_exceptions: pd.DataFrame | None = None,
	df_suggestions: pd.DataFrame | None = None,
	client: Any = None,
	share_public: bool | None = None,
	result_meta: dict[str, Any] | None = None,
	summarize_sections: Any = None,
) -> dict[str, Any]:
	"""Write all report sheets; callers may inject a gspread-compatible client."""
	service_email = None
	if client is None:
		client, service_email = _create_client()
	spreadsheet_key = config.SPREADSHEET_KEY
	if not spreadsheet_key:
		raise ValueError("SPREADSHEET_KEY belum dikonfigurasi di environment.")
	try:
		spreadsheet = _run_with_retry(lambda: client.open_by_key(spreadsheet_key))
	except Exception as error:
		_raise_google_error(error, service_email, client)

	result_payload = build_result_payload(df_result)
	meta = result_meta or result_payload["meta"]
	dashboard_payload = build_dashboard_payload(metrics, summary_text, df_result, meta)
	unmatched_payload = build_unmatched_payload(df_unmatched, df_debit_exceptions, df_suggestions, summarize_sections)
	payloads = (result_payload, dashboard_payload, unmatched_payload)
	worksheets = _ensure_worksheets(spreadsheet)
	for index, (title, payload) in enumerate(zip(_EXPECTED_SHEETS, payloads)):
		try:
			_write_payload(spreadsheet, worksheets[title], payload, index)
		except Exception as error:
			_raise_google_error(error, service_email, client)
	should_share = config.SHARE_PUBLIC if share_public is None else share_public
	if should_share:
		try:
			_run_with_retry(lambda: spreadsheet.share(email_address=None, perm_type="anyone", role="reader"))
		except Exception as error:
			_raise_google_error(error, service_email, client)
		logger.info("Google Sheet dibagikan untuk akses publik: %s", getattr(spreadsheet, "url", ""))
	return {
		"url": getattr(spreadsheet, "url", ""),
		"sheet_rows": {title: len(payload["values"]) for title, payload in zip(_EXPECTED_SHEETS, payloads)},
		"result_meta": meta,
		"shared_publicly": bool(should_share),
	}


def export_to_xlsx(
	df_result: pd.DataFrame,
	summary_text: str,
	metrics: dict[str, Any],
	df_unmatched: pd.DataFrame,
	*,
	path: str | Path,
	df_debit_exceptions: pd.DataFrame | None = None,
	df_suggestions: pd.DataFrame | None = None,
	summarize_sections: Any = None,
) -> dict[str, Any]:
	"""Write the existing three-sheet report payloads to a formatted XLSX file."""
	from openpyxl import Workbook
	from openpyxl.styles import Alignment, Font, PatternFill
	from openpyxl.utils import get_column_letter

	result_payload = build_result_payload(df_result)
	meta = result_payload["meta"]
	dashboard_payload = build_dashboard_payload(metrics, summary_text, df_result, meta)
	unmatched_payload = build_unmatched_payload(df_unmatched, df_debit_exceptions, df_suggestions, summarize_sections)
	workbook = Workbook()
	workbook.remove(workbook.active)
	for title, payload in zip(_EXPECTED_SHEETS, (result_payload, dashboard_payload, unmatched_payload)):
		worksheet = workbook.create_sheet(title)
		for row_index, row in enumerate(payload["values"], start=1):
			for column_index, value in enumerate(row, start=1):
				worksheet.cell(row_index, column_index, value)
		for merge in payload["merges"]:
			worksheet.merge_cells(merge["range"])
		for spec in payload["formats"]:
			for row in worksheet[spec["range"]]:
				for cell in row:
					if spec.get("background"):
						cell.fill = PatternFill("solid", fgColor=spec["background"].lstrip("#"))
					if spec.get("bold"):
						cell.font = Font(bold=True, color=spec.get("font_color", "#000000").lstrip("#"))
					elif spec.get("font_color"):
						cell.font = Font(color=spec["font_color"].lstrip("#"))
					if spec.get("wrap") or spec.get("align"):
						cell.alignment = Alignment(
							wrap_text=bool(spec.get("wrap")),
							horizontal=spec.get("align", "left").lower(),
						)
					if spec.get("number_format"):
						cell.number_format = spec["number_format"]
		for width in payload["column_widths"]:
			worksheet.column_dimensions[width["column"]].width = max(8, width["width"] / 7)
		worksheet.freeze_panes = f"{get_column_letter(payload['freeze']['columns'] + 1)}{payload['freeze']['rows'] + 1}"
		for height in payload.get("row_heights", []):
			worksheet.row_dimensions[height["row"]].height = height["height"]
	output_path = Path(path)
	output_path.parent.mkdir(parents=True, exist_ok=True)
	workbook.save(output_path)
	workbook.close()
	return {"path": str(output_path), "sheet_rows": {sheet.title: sheet.max_row for sheet in workbook.worksheets}, "result_meta": meta}