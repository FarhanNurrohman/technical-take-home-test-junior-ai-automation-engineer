"""Print raw workbook cells needed to verify the input layouts."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import xlrd
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
GL_PATH = DATA_DIR / "GL - Advances Other - April 2026.xls"
WP_PATH = DATA_DIR / "Working Paper Advances and Prepayment-Soal.xlsx"


def format_cell(value: Any, cell_type: str) -> str:
    """Return a printable representation including the Python cell type."""
    if isinstance(value, (date, datetime)):
        value_text = value.isoformat()
    else:
        value_text = repr(value)
    return f"{value_text} [python={type(value).__name__}, cell={cell_type}]"


def inspect_gl() -> None:
    workbook = xlrd.open_workbook(GL_PATH, formatting_info=False)
    print(f"GL: {GL_PATH.name}")
    for sheet in workbook.sheets():
        print(f"Sheet: {sheet.name!r} ({sheet.nrows} rows x {sheet.ncols} columns)")
        for row_index in range(min(15, sheet.nrows)):
            cells = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                    cell_type = "XL_CELL_DATE"
                else:
                    value = cell.value
                    cell_type = xlrd.sheet.ctype_text.get(cell.ctype, str(cell.ctype))
                cells.append(
                    f"{get_column_letter(column_index + 1)}{row_index + 1}="
                    f"{format_cell(value, cell_type)}"
                )
            print(f"  row {row_index + 1}: " + " | ".join(cells))


def inspect_working_paper() -> None:
    workbook = openpyxl.load_workbook(WP_PATH, data_only=False, read_only=False)
    print(f"Working Paper: {WP_PATH.name}")
    for sheet in workbook.worksheets:
        print(f"Sheet: {sheet.title!r} ({sheet.max_row} rows x {sheet.max_column} columns)")
        for row_index in range(1, min(25, sheet.max_row) + 1):
            cells = []
            for column_index in range(1, sheet.max_column + 1):
                cell = sheet.cell(row_index, column_index)
                if cell.value is not None:
                    cells.append(
                        f"{get_column_letter(column_index)}{row_index}="
                        f"{format_cell(cell.value, cell.data_type)}"
                    )
            print(f"  row {row_index}: " + (" | ".join(cells) if cells else "<empty>"))
    workbook.close()


def main() -> None:
    inspect_gl()
    inspect_working_paper()


if __name__ == "__main__":
    main()